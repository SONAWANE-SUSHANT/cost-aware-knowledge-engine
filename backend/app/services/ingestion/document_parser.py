from pathlib import Path

import fitz
from docx import Document
from pptx import Presentation


def extract_pdf(file_path: Path):
    pages = []

    document = fitz.open(file_path)

    for page_number, page in enumerate(document, start=1):
        text = page.get_text("text").strip()

        if text:
            pages.append({
                "page_number": page_number,
                "text": text
            })

    document.close()

    return pages


def extract_docx(file_path: Path):
    document = Document(file_path)

    text = []

    for paragraph in document.paragraphs:
        content = paragraph.text.strip()

        if content:
            text.append(content)

    return [
        {
            "page_number": None,
            "text": "\n".join(text)
        }
    ]


def extract_pptx(file_path: Path):
    presentation = Presentation(file_path)

    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        slide_text = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                content = shape.text.strip()

                if content:
                    slide_text.append(content)

        if slide_text:
            slides.append({
                "page_number": slide_number,
                "text": "\n".join(slide_text)
            })

    return slides


def extract_txt(file_path: Path):
    text = file_path.read_text(
        encoding="utf-8",
        errors="ignore"
    )

    return [
        {
            "page_number": None,
            "text": text.strip()
        }
    ]


def extract_text(file_path: Path):

    extension = file_path.suffix.lower()

    if extension == ".pdf":
        return extract_pdf(file_path)

    if extension == ".docx":
        return extract_docx(file_path)

    if extension == ".pptx":
        return extract_pptx(file_path)

    if extension == ".txt":
        return extract_txt(file_path)

    raise ValueError(
        f"Unsupported file type: {extension}"
    )